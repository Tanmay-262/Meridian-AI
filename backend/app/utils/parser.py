import os
from pypdf import PdfReader
import docx


def extract_text_from_pdf(file_path: str) -> str:
    """
    Extracts text page-by-page from a PDF file.
    If text streams are empty (scanned/image PDF), generates a structural metadata fallback.
    """
    reader = PdfReader(file_path)
    text_content = []

    for page in reader.pages:
        page_text = page.extract_text()
        if page_text and page_text.strip():
            text_content.append(page_text.strip())

    extracted = "\n\n".join(text_content)
    if extracted.strip():
        return extracted

    # Fallback for scanned / image-based PDFs
    filename = os.path.basename(file_path)
    num_pages = len(reader.pages)
    return (
        f"Document Title: {filename}\n"
        f"Document Type: Scanned Image PDF Document\n"
        f"Page Count: {num_pages} page(s)\n"
        f"Context: Scanned document file indexed by filename and structure."
    )


def extract_text_from_docx(file_path: str) -> str:
    """
    Extracts text paragraph-by-paragraph and from tables in a DOCX file.
    """
    doc = docx.Document(file_path)
    text_content = []

    for para in doc.paragraphs:
        if para.text.strip():
            text_content.append(para.text.strip())

    # Also extract text from tables inside the docx
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_content.append(" | ".join(row_text))

    extracted = "\n".join(text_content)
    if extracted.strip():
        return extracted

    # Fallback for empty/image DOCX files
    filename = os.path.basename(file_path)
    return (
        f"Document Title: {filename}\n"
        f"Document Type: Word Document (.docx)\n"
        f"Context: Document file indexed by filename and metadata."
    )


def extract_text_from_file(file_path: str) -> str:
    """
    Parses a file based on its extension and returns the extracted raw text.
    Throws ValueError for unsupported formats.
    """
    _, ext = os.path.splitext(file_path.lower())

    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_path)
    elif ext in [".txt", ".md", ".csv", ".json"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    else:
        raise ValueError(f"Unsupported file format: {ext}")
